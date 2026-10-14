from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy, copy
from lxml import etree as E
from openpyxl import load_workbook
import hashlib, json, re, posixpath
from io import BytesIO
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from config import load_config

def require(condition, message):
    if not condition:
        raise ValueError(message)
CONFIG, configured_path = load_config()
SOURCE = json.loads(configured_path('sourceData').read_text(encoding='utf-8'))
PLAN = json.loads(configured_path('assemblyPlan').read_text(encoding='utf-8'))
TARGET = configured_path('nmOutput', output=True)
S = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P = 'http://schemas.openxmlformats.org/package/2006/relationships'
C = 'http://schemas.openxmlformats.org/package/2006/content-types'
ns = {'s': S, 'r': R}
q = lambda n: f'{{{S}}}{n}'
sha = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
for path, value in SOURCE['protected_sha256'].items():
    require(sha(path) == value, path)

def read_parts(path):
    with ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}

def sheet_map(parts):
    rels = {x.get('Id'): x.get('Target') for x in E.fromstring(parts['xl/_rels/workbook.xml.rels'])}
    return {x.get('name'): posixpath.normpath('xl/' + rels[x.get(f'{{{R}}}id')]) if not rels[x.get(f'{{{R}}}id')].startswith('/') else rels[x.get(f'{{{R}}}id')].lstrip('/') for x in E.fromstring(parts['xl/workbook.xml']).find(q('sheets'))}

def xml(root):
    return E.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True)

def set_text(cell, value):
    for child in list(cell):
        if E.QName(child).localname in ('v', 'is', 'f'):
            cell.remove(child)
    cell.set('t', 'inlineStr')
    t = E.SubElement(E.SubElement(cell, q('is')), q('t'))
    t.text = value
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
parts = read_parts(configured_path('screeningBase'))
original = dict(parts)
donor = read_parts(configured_path('authoredChanges'))
paths = sheet_map(parts)
donor_paths = sheet_map(donor)
authored = load_workbook(configured_path('authoredChanges'), read_only=True)
for name, spec in PLAN['metadata'].items():
    rows = spec['values']
    actual = [[authored[name].cell(i + 1, j + 1).value for j in range(len(rows[0]))] for i in range(len(rows))]
    require(actual == rows, name)
    root = E.fromstring(parts[paths[name]])
    data = root.find(q('sheetData'))
    oldrows = {int(x.get('r')): x for x in data}
    authorrows = {int(x.get('r')): x for x in E.fromstring(donor[donor_paths[name]]).find(q('sheetData'))}
    for item in list(data):
        data.remove(item)
    for i, (origin, values) in enumerate(zip(spec['sourceRows'], rows), 1):
        row = deepcopy(oldrows[origin])
        row.set('r', str(i))
        for cell in row:
            cell.set('r', re.sub('\\d+$', str(i), cell.get('r')))
        for j, value in enumerate(values):
            address = chr(65 + j) + str(i)
            cell = row.find(f"s:c[@r='{address}']", ns)
            if cell is None:
                cell = E.SubElement(row, q('c'), r=address)
            set_text(cell, value)
        for key in ('ht', 'customHeight'):
            if authorrows.get(i) is not None and authorrows[i].get(key) is not None:
                row.set(key, authorrows[i].get(key))
        data.append(row)
    dim = root.find(q('dimension'))
    if dim is not None:
        dim.set('ref', f'A1:{chr(64 + len(rows[0]))}{len(rows)}')
    parts[paths[name]] = xml(root)
styles = E.fromstring(parts['xl/styles.xml'])
donor_styles = E.fromstring(donor['xl/styles.xml'])

def col(root, tag):
    el = root.find(q(tag))
    if el is None:
        el = E.SubElement(root, q(tag), count='0')
    return el
numfmts = col(styles, 'numFmts')
known = {x.get('formatCode'): int(x.get('numFmtId')) for x in numfmts}
nextid = max([163] + list(known.values())) + 1
fmtmap = {}
for el in col(donor_styles, 'numFmts'):
    code = el.get('formatCode')
    old = int(el.get('numFmtId'))
    if code not in known:
        known[code] = nextid
        nextid += 1
        item = deepcopy(el)
        item.set('numFmtId', str(known[code]))
        numfmts.append(item)
    fmtmap[old] = known[code]
numfmts.set('count', str(len(numfmts)))
maps = {}
for tag in ('fonts', 'fills', 'borders'):
    dest = col(styles, tag)
    src = col(donor_styles, tag)
    offset = len(dest)
    maps[tag] = {i: offset + i for i in range(len(src))}
    for item in src:
        dest.append(deepcopy(item))
    dest.set('count', str(len(dest)))

def remap_xf(el, with_xfid=False):
    for attr, tag in [('fontId', 'fonts'), ('fillId', 'fills'), ('borderId', 'borders')]:
        if el.get(attr) is not None:
            el.set(attr, str(maps[tag][int(el.get(attr))]))
    if el.get('numFmtId') is not None:
        n = int(el.get('numFmtId'))
        el.set('numFmtId', str(fmtmap.get(n, n)))
    if with_xfid and el.get('xfId') is not None:
        el.set('xfId', str(maps['cellStyleXfs'][int(el.get('xfId'))]))
for tag in ('cellStyleXfs', 'cellXfs'):
    dest = col(styles, tag)
    src = col(donor_styles, tag)
    offset = len(dest)
    maps[tag] = {i: offset + i for i in range(len(src))}
    for el in src:
        item = deepcopy(el)
        remap_xf(item, tag == 'cellXfs')
        dest.append(item)
    dest.set('count', str(len(dest)))
parts['xl/styles.xml'] = xml(styles)
detail = E.fromstring(donor[donor_paths[PLAN['sheetName']]])
shared = E.fromstring(donor['xl/sharedStrings.xml']) if 'xl/sharedStrings.xml' in donor else None
for cell in detail.findall('.//s:sheetData/s:row/s:c', ns):
    if cell.get('s') is not None:
        cell.set('s', str(maps['cellXfs'][int(cell.get('s'))]))
    if cell.get('t') == 's':
        value = cell.find(q('v'))
        content = shared[int(value.text)]
        cell.remove(value)
        cell.set('t', 'inlineStr')
        inline = E.SubElement(cell, q('is'))
        for child in content:
            inline.append(deepcopy(child))
for row in detail.findall('.//s:row', ns):
    if row.get('s') is not None:
        row.set('s', str(maps['cellXfs'][int(row.get('s'))]))
for column in detail.findall('.//s:col', ns):
    if column.get('style') is not None:
        column.set('style', str(maps['cellXfs'][int(column.get('style'))]))
detailpath = 'xl/worksheets/sheet39.xml'
require(detailpath not in parts, 'Validation failed')
parts[detailpath] = xml(detail)
wb = E.fromstring(parts['xl/workbook.xml'])
sheets = wb.find(q('sheets'))
rels = E.fromstring(parts['xl/_rels/workbook.xml.rels'])
rid = 'rIdNmS1ScreeningDetail'
require(rid not in {x.get('Id') for x in rels}, 'Validation failed')
E.SubElement(rels, f'{{{P}}}Relationship', Id=rid, Type=R + '/worksheet', Target='worksheets/sheet39.xml')
item = E.SubElement(sheets, q('sheet'), name=PLAN['sheetName'], sheetId='39')
item.set(f'{{{R}}}id', rid)
lookup = {s.get('name'): s for s in sheets}
require(set(lookup) == set(PLAN['order']), 'Validation failed')
for item in list(sheets):
    sheets.remove(item)
for name in PLAN['order']:
    sheets.append(lookup[name])
parts['xl/workbook.xml'] = xml(wb)
parts['xl/_rels/workbook.xml.rels'] = xml(rels)
ct = E.fromstring(parts['[Content_Types].xml'])
E.SubElement(ct, f'{{{C}}}Override', PartName='/' + detailpath, ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml')
parts['[Content_Types].xml'] = xml(ct)
if 'docProps/app.xml' in parts:
    app = E.fromstring(parts['docProps/app.xml'])
    vt = 'http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes'
    for el in app:
        if E.QName(el).localname == 'TitlesOfParts':
            vec = el[0]
            for child in list(vec):
                vec.remove(child)
            vec.set('size', '39')
            for name in PLAN['order']:
                E.SubElement(vec, f'{{{vt}}}lpstr').text = name
        if E.QName(el).localname == 'HeadingPairs':
            items = list(el[0])
            for i in range(len(items) - 1):
                if ''.join(items[i].itertext()) == 'Worksheets':
                    items[i + 1][0].text = '39'
    parts['docProps/app.xml'] = xml(app)
if CONFIG.get('files', {}).get('layoutReference'):
    layout = read_parts(configured_path('layoutReference'))
    layout_paths = sheet_map(layout)
    for name in [*PLAN['metadata'], PLAN['sheetName']]:
        destination = paths.get(name, detailpath)
        values_root = E.fromstring(parts[destination])
        layout_root = E.fromstring(layout[layout_paths[name]])
        values_cells = {c.get('r'): c for c in values_root.findall('.//s:sheetData/s:row/s:c', ns)}
        layout_cells = {c.get('r'): c for c in layout_root.findall('.//s:sheetData/s:row/s:c', ns)}
        require(set(values_cells).issubset(layout_cells), name)
        for address, cell in layout_cells.items():
            for child in list(cell):
                if E.QName(child).localname in ('v', 'is', 'f'):
                    cell.remove(child)
            cell.attrib.pop('t', None)
            value = values_cells.get(address)
            if value is not None:
                if value.get('t') is not None:
                    cell.set('t', value.get('t'))
                for child in value:
                    if E.QName(child).localname in ('v', 'is', 'f'):
                        cell.append(deepcopy(child))
        parts[destination] = xml(layout_root)

    def signature(element):
        if element is None:
            return None
        return (element.tag, tuple(sorted(element.attrib.items())), (element.text or '').strip(), tuple((signature(child) for child in element)))
    base_styles = E.fromstring(original['xl/styles.xml'])
    layout_styles = E.fromstring(layout['xl/styles.xml'])
    for tag in ('dxfs', 'cellStyles', 'tableStyles', 'colors'):
        require(signature(base_styles.find(q(tag))) == signature(layout_styles.find(q(tag))), tag)
    def custom_formats(styles):
        element = styles.find(q('numFmts'))
        formats = list(element) if element is not None else []
        mapping = {item.get('numFmtId'): item.get('formatCode') for item in formats}
        require(len(mapping) == len(formats), 'Duplicate number format identifiers')
        return mapping
    original_formats = custom_formats(base_styles)
    reference_formats = custom_formats(layout_styles)
    require(all(reference_formats.get(key) == value for key, value in original_formats.items()), 'Number format mapping changed')
    for tag in ('fonts', 'fills', 'borders', 'cellStyleXfs', 'cellXfs'):
        first = base_styles.find(q(tag))
        second = layout_styles.find(q(tag))
        if first is not None:
            require(second is not None and len(second) >= len(first), tag)
            require(all((signature(a) == signature(b) for a, b in zip(first, second))), tag)
    parts['xl/styles.xml'] = layout['xl/styles.xml']
buffer = BytesIO()
with ZipFile(buffer, 'w', compression=ZIP_DEFLATED) as z:
    for name, data in parts.items():
        z.writestr(name, data)
buffer.seek(0)
old = load_workbook(configured_path('screeningBase'), data_only=False)
final = load_workbook(buffer, data_only=False)
require(final.sheetnames == PLAN['order'], 'Validation failed')
unchanged = []
compared = 0
style_pairs = set()
for name in old.sheetnames:
    if name in PLAN['metadata']:
        continue
    require(parts[paths[name]] == original[paths[name]], name)
    unchanged.append(name)
    a, b = (old[name], final[name])
    require((a.max_row, a.max_column) == (b.max_row, b.max_column), 'Validation failed')
    for row in a:
        for cell in row:
            other = b[cell.coordinate]
            require((cell.value, cell.data_type, cell.number_format) == (other.value, other.data_type, other.number_format), (name, cell.coordinate))
            style_pair = (cell.style_id, other.style_id)
            if style_pair not in style_pairs:
                for attribute in ('font', 'fill', 'border', 'alignment', 'protection'):
                    require(copy(getattr(cell, attribute)) == copy(getattr(other, attribute)), (name, cell.coordinate, attribute))
                style_pairs.add(style_pair)
            compared += 1
d = final[PLAN['sheetName']]
require(d['A1'].value == SOURCE['title'], 'Validation failed')
for i, row in enumerate(SOURCE['typed_rows'], 4):
    for j, value in enumerate(row, 1):
        require(d.cell(i, j).value == value, (i, j))
for i, note in enumerate(SOURCE['notes'], 37):
    require(d.cell(i, 1).value == note, i)
for path, value in SOURCE['protected_sha256'].items():
    require(sha(path) == value, path)
final.close()
old.close()
authored.close()
TARGET.write_bytes(buffer.getvalue())
