import hashlib
import json
import re
import sys
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from config import load_config


config, path = load_config()
source = path('screeningDocument')
base = path('screeningBase')
document = Document(source)
paragraphs = [p.text for p in document.paragraphs if p.text]
rows = [[cell.text for cell in row.cells] for row in document.tables[0].rows]
typed = [[int(value.replace(',', ''))
          if column in (2, 3, 4) and re.fullmatch(r'[\d,]+', value)
          else (value or None)
          for column, value in enumerate(row)] for row in rows]
workbook = load_workbook(base, read_only=True)
metadata = {name: [list(row) for row in workbook[name].iter_rows(values_only=True)]
            for name in ['Contents', 'Data_dictionary', 'Notes']}
payload = {
    'source': str(source),
    'base': str(base),
    'title': paragraphs[0],
    'notes': paragraphs[1:],
    'source_rows': rows,
    'typed_rows': typed,
    'metadata': metadata,
    'base_order': workbook.sheetnames,
    'protected_sha256': {str(item): hashlib.sha256(item.read_bytes()).hexdigest()
                         for item in [source, base]},
}
workbook.close()
path('sourceData', output=True).write_text(
    json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
